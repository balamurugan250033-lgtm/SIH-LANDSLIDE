from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).parent
OUTPUT = ROOT / "SIH2026-IDEA-Presentation-Landslide-NER-ENHANCED.pptx"

W = Inches(13.333)
H = Inches(7.5)
NAVY = RGBColor(10, 18, 38)
INK = RGBColor(19, 30, 51)
MUTED = RGBColor(102, 116, 139)
WHITE = RGBColor(248, 250, 252)
SLATE = RGBColor(226, 232, 240)
BLUE = RGBColor(37, 99, 235)
CYAN = RGBColor(34, 211, 238)
ORANGE = RGBColor(249, 115, 22)
AMBER = RGBColor(245, 158, 11)
GREEN = RGBColor(22, 163, 74)
RED = RGBColor(220, 38, 38)
PALE_BLUE = RGBColor(239, 246, 255)
PALE_ORANGE = RGBColor(255, 247, 237)
PALE_GREEN = RGBColor(240, 253, 244)


def rgb(hex_value):
    return RGBColor.from_string(hex_value.replace("#", "").upper())


def add_box(slide, x, y, w, h, fill, line=None, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    shape.line.width = Pt(0.8)
    return shape


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.05, break_line=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    if break_line:
        lines = text.split("\n")
        for index, line in enumerate(lines):
            if index:
                p = tf.add_paragraph()
                p.alignment = align
            run = p.add_run()
            run.text = line
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    else:
        run = p.add_run()
        run.text = text
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_rich_text(slide, runs, x, y, w, h, size=18, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.03)
    p = tf.paragraphs[0]
    for text, bold, run_color in runs:
        run = p.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = run_color or color
    return box


def add_line(slide, x1, y1, x2, y2, color=SLATE, width=1.3):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_footer(slide, number, dark=False):
    color = rgb("CBD5E1") if dark else MUTED
    add_text(slide, "RIT Chennai  |  Smart India Hackathon 2026", 0.55, 7.12, 5.2, 0.18, 8.5, color)
    add_text(slide, f"{number:02d} / 06", 12.1, 7.12, 0.68, 0.18, 8.5, color, bold=True, align=PP_ALIGN.RIGHT)


def add_top_label(slide, kicker, title, number, dark=False):
    add_text(slide, kicker.upper(), 0.58, 0.34, 4.3, 0.25, 9, CYAN if dark else BLUE, bold=True)
    add_text(slide, title, 0.55, 0.68, 10.7, 0.65, 27, WHITE if dark else INK, bold=True)
    add_footer(slide, number, dark)


def add_chip(slide, text, x, y, w, fill, color=INK):
    add_box(slide, x, y, w, 0.36, fill, fill)
    add_text(slide, text, x + 0.08, y + 0.075, w - 0.16, 0.2, 10, color, bold=True, align=PP_ALIGN.CENTER)


def add_metric(slide, value, label, x, y, w, accent=BLUE, dark=False):
    fill = rgb("162441") if dark else WHITE
    add_box(slide, x, y, w, 0.92, fill, rgb("2A3C61") if dark else SLATE)
    add_text(slide, value, x + 0.16, y + 0.12, w - 0.3, 0.34, 23, accent if dark else INK, bold=True)
    add_text(slide, label, x + 0.16, y + 0.57, w - 0.3, 0.18, 9.5, rgb("AAB8D0") if dark else MUTED)


def add_bullet(slide, text, x, y, w, color=INK, accent=BLUE, size=13):
    add_box(slide, x, y + 0.08, 0.11, 0.11, accent, accent, radius=False)
    add_text(slide, text, x + 0.22, y, w - 0.22, 0.42, size, color)


def add_screenshot(slide, path, x, y, w, h, label, dark=False):
    frame = add_box(slide, x - 0.08, y - 0.08, w + 0.16, h + 0.16, rgb("25334F") if dark else WHITE, rgb("334A73") if dark else SLATE)
    frame.shadow.inherit = False
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    add_box(slide, x + 0.14, y + h - 0.36, min(2.15, w - 0.28), 0.25, NAVY, NAVY)
    add_text(slide, label.upper(), x + 0.21, y + h - 0.31, min(2.0, w - 0.4), 0.15, 7.5, WHITE, bold=True)


def setup_slide(prs, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY if dark else WHITE
    return slide


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1. Cover
    slide = setup_slide(prs, dark=True)
    add_box(slide, 0, 0, 13.333, 0.12, ORANGE, ORANGE, radius=False)
    add_text(slide, "SMART INDIA HACKATHON 2026", 0.62, 0.55, 4.6, 0.28, 10, CYAN, bold=True)
    add_text(slide, "Landslide Sentinel", 0.6, 1.27, 7.4, 0.74, 38, WHITE, bold=True)
    add_text(slide, "AI-powered early warning for the North-Eastern Region", 0.64, 2.16, 6.4, 0.5, 18, rgb("D5E1F5"))
    add_text(slide, "From changing ground conditions to clear, local action.", 0.64, 2.78, 5.9, 0.32, 12, rgb("AAB8D0"))
    add_chip(slide, "AI RISK SCORING", 0.64, 3.58, 1.55, rgb("20345B"), CYAN)
    add_chip(slide, "OFFLINE-FIRST", 2.32, 3.58, 1.42, rgb("20345B"), CYAN)
    add_chip(slide, "MESH RELAY", 3.87, 3.58, 1.24, rgb("20345B"), CYAN)
    add_metric(slide, "9", "monitored regions", 0.64, 5.12, 1.65, CYAN, True)
    add_metric(slide, "4-tier", "risk classification", 2.48, 5.12, 1.85, ORANGE, True)
    add_metric(slide, "3", "connected interfaces", 4.52, 5.12, 1.72, GREEN, True)
    add_box(slide, 8.0, 1.02, 4.48, 5.5, rgb("101D3B"), rgb("29416C"))
    add_text(slide, "EARLY WARNING LOOP", 8.42, 1.38, 3.5, 0.25, 10, CYAN, bold=True)
    steps = [("01", "Observe", "rainfall  |  soil  |  slope", BLUE), ("02", "Score", "XGBoost risk inference", ORANGE), ("03", "Alert", "web  |  SMS  |  push", RED), ("04", "Act", "report  |  route  |  evacuate", GREEN)]
    for idx, (num, title, detail, accent) in enumerate(steps):
        y = 1.94 + idx * 1.02
        add_box(slide, 8.42, y, 0.55, 0.55, accent, accent)
        add_text(slide, num, 8.42, y + 0.16, 0.55, 0.16, 9, WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, title, 9.23, y - 0.01, 2.2, 0.27, 17, WHITE, bold=True)
        add_text(slide, detail, 9.23, y + 0.3, 2.75, 0.2, 10, rgb("AAB8D0"))
        if idx < len(steps) - 1:
            add_line(slide, 8.69, y + 0.58, 8.69, y + 0.96, rgb("40577E"), 1.6)
    add_text(slide, "Team RIT Chennai  |  Problem Statement: AI-Powered Landslide Early Warning System for NER", 0.62, 6.88, 8.5, 0.22, 8.5, rgb("AAB8D0"))
    add_text(slide, "01 / 06", 12.1, 7.12, 0.68, 0.18, 8.5, rgb("CBD5E1"), bold=True, align=PP_ALIGN.RIGHT)

    # 2. Problem and solution
    slide = setup_slide(prs)
    add_top_label(slide, "The opportunity", "Make a warning useful before the slope moves.", 2)
    add_text(slide, "Landslide risk is not only a prediction problem. It is a last-mile coordination problem.", 0.58, 1.47, 9.4, 0.35, 14, MUTED)
    columns = [("FRAGILE SIGNALS", "Rainfall, soil saturation, slope angle and vibration can change faster than manual reporting.", ORANGE, PALE_ORANGE), ("FRAGMENTED ACTION", "Citizens, administrators and responders need the same local picture, in the same moment.", BLUE, PALE_BLUE), ("LOW CONNECTIVITY", "An internet outage should not erase the last alert or block a new hazard report.", GREEN, PALE_GREEN)]
    for idx, (head, body, accent, fill) in enumerate(columns):
        x = 0.58 + idx * 4.22
        add_box(slide, x, 2.08, 3.86, 1.55, fill, fill)
        add_box(slide, x, 2.08, 0.1, 1.55, accent, accent, radius=False)
        add_text(slide, head, x + 0.25, 2.3, 3.2, 0.25, 11, accent, bold=True)
        add_text(slide, body, x + 0.25, 2.72, 3.18, 0.62, 12, INK)
    add_text(slide, "ONE SHARED OPERATING PICTURE", 0.58, 4.18, 4.2, 0.22, 10, BLUE, bold=True)
    flow = [("Sense", "Live observations"), ("Infer", "Explainable risk tier"), ("Decide", "Admin + citizen context"), ("Reach", "Online or mesh")]
    for idx, (title, detail) in enumerate(flow):
        x = 0.58 + idx * 3.1
        add_box(slide, x, 4.67, 2.5, 1.1, NAVY, NAVY)
        add_text(slide, f"0{idx + 1}", x + 0.18, 4.85, 0.38, 0.2, 10, CYAN, bold=True)
        add_text(slide, title, x + 0.66, 4.82, 1.5, 0.24, 16, WHITE, bold=True)
        add_text(slide, detail, x + 0.66, 5.16, 1.62, 0.2, 9.5, rgb("AAB8D0"))
        if idx < 3:
            add_line(slide, x + 2.54, 5.22, x + 3.0, 5.22, ORANGE, 2.5)
    add_text(slide, "Outcome: a risk signal becomes a visible, actionable safety decision.", 0.58, 6.35, 8.8, 0.28, 14, INK, bold=True)
    add_footer(slide, 2)

    # 3. Technical approach
    slide = setup_slide(prs, dark=True)
    add_top_label(slide, "Technical approach", "A resilient path from signal to safety action.", 3, dark=True)
    add_text(slide, "Modular services keep live monitoring, decision logic and last-mile delivery independently testable.", 0.58, 1.47, 8.7, 0.32, 13, rgb("AAB8D0"))
    layers = [("DATA", "sensor feeds  |  citizen reports", CYAN), ("INTELLIGENCE", "XGBoost inference  |  4-tier decision", ORANGE), ("OPERATIONS", "FastAPI  |  SQLite  |  alert dispatch", BLUE), ("LAST MILE", "web  |  mobile  |  mesh relay", GREEN)]
    for idx, (head, body, accent) in enumerate(layers):
        y = 2.14 + idx * 0.77
        add_box(slide, 0.58, y, 5.75, 0.58, rgb("162441"), rgb("2A3C61"))
        add_box(slide, 0.58, y, 0.11, 0.58, accent, accent, radius=False)
        add_text(slide, head, 0.88, y + 0.14, 1.32, 0.18, 10, accent, bold=True)
        add_text(slide, body, 2.28, y + 0.13, 3.7, 0.2, 12, WHITE, bold=True)
        if idx < 3:
            add_line(slide, 3.45, y + 0.6, 3.45, y + 0.75, rgb("40577E"), 1.5)
    add_text(slide, "Proof in the product", 7.1, 1.92, 2.6, 0.26, 11, CYAN, bold=True)
    add_screenshot(slide, ROOT / "docs" / "screenshots" / "admin-dashboard.png", 7.05, 2.3, 5.38, 3.36, "admin command view", dark=True)
    add_chip(slide, "ROLE-BASED", 7.05, 5.95, 1.14, rgb("20345B"), CYAN)
    add_chip(slide, "GIS-AWARE", 8.35, 5.95, 1.12, rgb("20345B"), CYAN)
    add_chip(slide, "AUDITABLE", 9.64, 5.95, 1.12, rgb("20345B"), CYAN)
    add_footer(slide, 3, dark=True)

    # 4. Feasibility
    slide = setup_slide(prs)
    add_top_label(slide, "Feasibility and viability", "Designed for the conditions it must survive.", 4)
    add_text(slide, "The architecture is practical today and extensible tomorrow: familiar web services, local persistence and a tested offline path.", 0.58, 1.47, 10.5, 0.35, 13, MUTED)
    add_screenshot(slide, ROOT / "docs" / "screenshots" / "citizen-map.png", 0.58, 2.07, 5.25, 2.95, "citizen GIS view")
    add_text(slide, "Offline-first behavior", 6.35, 2.08, 3.1, 0.26, 16, INK, bold=True)
    offline = [("CACHE", "Keep the last known region, alert and road status locally."), ("QUEUE", "Store hazard reports while the network is unavailable."), ("SYNC", "Replay queued reports automatically when connectivity returns."), ("RELAY", "Pass unseen alerts phone-to-phone with TTL-based deduplication.")]
    for idx, (head, body) in enumerate(offline):
        y = 2.55 + idx * 0.69
        add_box(slide, 6.35, y, 5.95, 0.51, PALE_BLUE if idx % 2 == 0 else PALE_GREEN, PALE_BLUE if idx % 2 == 0 else PALE_GREEN)
        add_text(slide, head, 6.55, y + 0.14, 0.72, 0.16, 9.5, BLUE if idx < 3 else GREEN, bold=True)
        add_text(slide, body, 7.48, y + 0.1, 4.45, 0.26, 10.5, INK)
    add_metric(slide, "6 suites", "backend + frontend + offline workflows", 0.58, 5.62, 2.3, BLUE)
    add_metric(slide, "SQLite", "local persistence for low-connectivity use", 3.06, 5.62, 2.55, ORANGE)
    add_metric(slide, "API-first", "ready for sensors, SMS and satellite feeds", 5.78, 5.62, 2.55, GREEN)
    add_metric(slide, "Pilotable", "deploy backend, web and APK independently", 8.5, 5.62, 2.55, CYAN)
    add_footer(slide, 4)

    # 5. Impact
    slide = setup_slide(prs, dark=True)
    add_top_label(slide, "Impact and benefits", "Give every role the next useful action.", 5, dark=True)
    add_text(slide, "The system turns uncertainty into a shared response: observe clearly, communicate early, move deliberately.", 0.58, 1.47, 8.7, 0.32, 13, rgb("AAB8D0"))
    add_screenshot(slide, ROOT / "docs" / "screenshots" / "citizen-platform.png", 0.58, 2.06, 4.35, 4.28, "citizen monitoring view", dark=True)
    outcomes = [("CITIZENS", "See local risk, alerts, safe routes and reporting tools in one place.", CYAN), ("AUTHORITIES", "Prioritize regions, issue targeted alerts and track road conditions.", ORANGE), ("RESPONDERS", "Carry the last known alert beyond the edge of connectivity.", GREEN)]
    for idx, (head, body, accent) in enumerate(outcomes):
        y = 2.14 + idx * 1.21
        add_box(slide, 5.45, y, 6.85, 0.94, rgb("162441"), rgb("2A3C61"))
        add_box(slide, 5.45, y, 0.1, 0.94, accent, accent, radius=False)
        add_text(slide, head, 5.78, y + 0.18, 1.7, 0.2, 10, accent, bold=True)
        add_text(slide, body, 7.52, y + 0.15, 4.38, 0.47, 12, WHITE)
    add_text(slide, "Success looks like", 5.45, 5.95, 2.0, 0.22, 10, CYAN, bold=True)
    add_text(slide, "a warning that still reaches people when the network does not.", 7.1, 5.9, 4.65, 0.32, 16, WHITE, bold=True)
    add_footer(slide, 5, dark=True)

    # 6. Research, validation and close
    slide = setup_slide(prs)
    add_top_label(slide, "Validation and next step", "Built to be tested in the field, then extended.", 6)
    add_text(slide, "Current implementation is validated across the decision path; the next milestone is real-world sensor and community pilot data.", 0.58, 1.47, 10.5, 0.35, 13, MUTED)
    add_text(slide, "VALIDATED NOW", 0.58, 2.1, 2.3, 0.22, 10, BLUE, bold=True)
    checks = ["Risk decision and threshold mapping", "Sensor ingestion and validation", "Citizen report submission", "Offline queue and reconnect sync", "Push / notification behavior", "Admin alert dispatch workflow"]
    for idx, item in enumerate(checks):
        x = 0.58 + (idx % 2) * 3.0
        y = 2.54 + (idx // 2) * 0.62
        add_box(slide, x, y, 2.7, 0.38, PALE_GREEN, PALE_GREEN)
        add_text(slide, "✓", x + 0.12, y + 0.09, 0.2, 0.16, 11, GREEN, bold=True)
        add_text(slide, item, x + 0.43, y + 0.08, 2.1, 0.2, 9.5, INK)
    add_box(slide, 6.95, 2.06, 5.35, 2.7, NAVY, NAVY)
    add_text(slide, "NEXT PILOT", 7.35, 2.44, 2.0, 0.22, 10, CYAN, bold=True)
    add_text(slide, "Connect live observation to local action.", 7.35, 2.82, 4.2, 0.4, 20, WHITE, bold=True)
    add_text(slide, "IoT sensors  →  district pilot  →  multilingual delivery  →  community feedback", 7.35, 3.52, 4.25, 0.46, 11, rgb("D5E1F5"))
    add_chip(slide, "SCALABLE", 7.35, 4.15, 1.04, rgb("20345B"), CYAN)
    add_chip(slide, "MEASURABLE", 8.55, 4.15, 1.18, rgb("20345B"), CYAN)
    add_chip(slide, "COMMUNITY-READY", 9.88, 4.15, 1.58, rgb("20345B"), CYAN)
    add_text(slide, "References", 0.58, 5.32, 1.5, 0.22, 11, BLUE, bold=True)
    add_text(slide, "FastAPI, SQLAlchemy, XGBoost, React/Vite, React Native, Expo, SQLite, Google Nearby Connections", 0.58, 5.68, 7.9, 0.28, 11, INK)
    add_text(slide, "Thank you", 9.1, 5.55, 2.7, 0.5, 27, INK, bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, "RIT Chennai  |  NER Landslide Early Warning System", 7.85, 6.15, 4.45, 0.22, 10, MUTED, align=PP_ALIGN.RIGHT)
    add_footer(slide, 6)

    prs.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()